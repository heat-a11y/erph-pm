#!/usr/bin/env python3
"""Build clean curriculum data from RPT files with better extraction."""
import os, re, json, sys
from datetime import datetime

RPT_DIR = "/home/home/Downloads/dskp&rpt/2_RPT_2026-20260717T123240Z-1-001/2_RPT_2026"
DSKP_DIR = "/home/home/Downloads/dskp&rpt/1_DSKP & RPT_2026-20260717T122750Z-1-001"

SUBJECT_INFO = {
    'BM': {'name': 'Bahasa Melayu', 'code': 'BM', 'color': '#1E3A5F'},
    'BI': {'name': 'Bahasa Inggeris', 'code': 'BI', 'color': '#4F7942'},
    'BC': {'name': 'Bahasa Cina', 'code': 'BC', 'color': '#C4956A'},
    'MT': {'name': 'Matematik', 'code': 'MT', 'color': '#8B5A2B'},
    'SN': {'name': 'Sains', 'code': 'SN', 'color': '#0D9488'},
    'PM': {'name': 'Pendidikan Moral', 'code': 'PM', 'color': '#7C3AED'},
    'PI': {'name': 'Pendidikan Islam', 'code': 'PI', 'color': '#059669'},
    'PJ': {'name': 'Pendidikan Jasmani', 'code': 'PJ', 'color': '#DC2626'},
    'PK': {'name': 'Pendidikan Kesihatan', 'code': 'PK', 'color': '#E11D48'},
    'PSV': {'name': 'Pendidikan Seni Visual', 'code': 'PSV', 'color': '#D97706'},
    'MZ': {'name': 'Pendidikan Muzik', 'code': 'MZ', 'color': '#7C3AED'},
    'SJ': {'name': 'Sejarah', 'code': 'SJ', 'color': '#B45309'},
    'RBT': {'name': 'Reka Bentuk dan Teknologi', 'code': 'RBT', 'color': '#6366F1'},
    'BA': {'name': 'Bahasa Arab', 'code': 'BA', 'color': '#0891B2'},
}

SUBJECT_NUM_MAP = {
    '1': 'BM', '2': 'BI', '3': 'BC', '4': 'MT', '5': 'SN',
    '6': 'PM', '6a': 'PM', '6b': 'PI', '6c': 'BA',
    '7': 'PJ', '7a': 'PJ', '7b': 'PK',
    '8': 'PSV', '8a': 'PSV', '8b': 'MZ',
    '9': 'SJ', '10': 'RBT',
    '11': 'PSV', '12': 'RBT',
}

def clean(s):
    if not s: return ''
    return str(s).strip().replace('\xa0', ' ').replace('\u200b', '').replace('\n', ' | ')

def detect_subj(filename):
    f = filename.lower().replace(' ', '_')
    for num, code in SUBJECT_NUM_MAP.items():
        t = f
        if t.startswith(f'{num}.' ) or t.startswith(f'{num}_') or t.startswith(num):
            return code
        if f' {num}.' in f' {t}': return code
    kw = {'bm': 'BM', 'bahasa_melayu': 'BM', 'bi': 'BI', 'bahasa_ingeris': 'BI',
          'bc': 'BC', 'bahasa_cina': 'BC', 'mt': 'MT', 'matematik': 'MT',
          'sn': 'SN', 'sains': 'SN', 'pm': 'PM', 'moral': 'PM',
          'pi': 'PI', 'pendidikan_islam': 'PI', 'pj': 'PJ', 'pk': 'PK',
          'psv': 'PSV', 'seni': 'PSV', 'mz': 'MZ', 'muzik': 'MZ',
          'sj': 'SJ', 'sejarah': 'SJ', 'rbt': 'RBT'}
    for k, v in kw.items():
        if k in f: return v
    return None

def get_year_no(dirname):
    m = re.search(r'(\d)', dirname)
    return int(m.group(1)) if m else 0

def extract_table_rows(table):
    """Extract all non-empty rows from a DOCX table."""
    import docx
    rows = []
    for r in table.rows:
        cells = [clean(c.text) for c in r.cells]
        if any(c for c in cells):
            rows.append(cells)
    return rows

def process_bm_table(rows):
    """Extract unit info from BM table."""
    units = []
    for r in rows:
        if len(r) < 3: continue
        minggu = r[0] if r[0] else ''
        tema = r[1] if len(r) > 1 and r[1] else ''
        sk = r[2] if len(r) > 2 and r[2] else ''
        sp = r[3] if len(r) > 3 and r[3] else ''
        catatan = r[4] if len(r) > 4 and r[4] else ''
        
        if not tema and not minggu: continue
        tema_lower = tema.lower()
        if 'pkjr' in tema_lower or 'orientasi' in tema_lower or 'transisi' in tema_lower:
            continue
        if 'cuti' in tema_lower or 'cuti' in minggu.lower(): continue
        
        # Extract clean title
        title = tema.replace('TEMA', '').replace('Unit', '').strip()
        title = re.sub(r'\s+', ' ', title).strip(' |')
        if not title or len(title) < 3: continue
        
        # Clean up SK
        sk_clean = clean(sk)
        sp_clean = clean(sp)
        
        # Parse minggu numbers
        weeks = re.findall(r'\d+', minggu)
        
        units.append({
            'minggu': minggu[:30],
            'title': title[:120],
            'sk': sk_clean[:200],
            'sp': sp_clean[:200],
        })
    return units

def process_mt_table(rows):
    """Extract from 3-column Chinese math/science table."""
    units = []
    header_skipped = False
    for r in rows:
        if not header_skipped and ('周' in r[0] or 'minggu' in r[0].lower()):
            header_skipped = True
            continue
        if len(r) < 2: continue
        
        minggu = clean(r[0]) if r[0] else ''
        topic = clean(r[2]) if len(r) > 2 and r[2] else (clean(r[1]) if len(r) > 1 and r[1] else '')
        
        t_lower = topic.lower()
        if not topic or len(topic) < 3: continue
        if 'cuti' in t_lower or 'program' in t_lower or 'orientasi' in t_lower:
            continue
        if '开学' in topic or '迎新' in topic: continue
        
        # Clean the title
        title = re.sub(r'\s*\d+\s*$', '', topic).strip()
        title = re.sub(r'\s+', ' ', title)[:120]
        
        units.append({
            'minggu': minggu[:30],
            'title': title,
        })
    return units

def process_bi_table(rows):
    """Extract from BI table."""
    units = []
    seen = set()
    for r in rows:
        if len(r) < 5: continue
        week = clean(r[0]) if r[0] else ''
        theme_topic = clean(r[1]) if len(r) > 1 and r[1] else ''
        lesson = clean(r[2]) if len(r) > 2 and r[2] else ''
        
        if not theme_topic: continue
        if 'transition' in theme_topic.lower(): continue
        
        # Extract meaningful theme/topic
        parts = theme_topic.split('|')
        title = parts[0].strip() if parts else theme_topic[:80]
        title = re.sub(r'\s+', ' ', title).strip()
        if not title or len(title) < 5: continue
        if title in seen: continue
        seen.add(title)
        
        units.append({
            'minggu': week[:25],
            'title': title[:120],
            'lesson': lesson[:50],
        })
    return units

def extract_all():
    data = {}  # key: f"{year}|{subj}"
    counts = {}
    
    import docx
    import openpyxl
    
    for year_dir in sorted(os.listdir(RPT_DIR)):
        year_path = os.path.join(RPT_DIR, year_dir)
        if not os.path.isdir(year_path): continue
        year_no = get_year_no(year_dir)
        if not year_no: continue
        year = f'Year {year_no}'
        
        for fname in sorted(os.listdir(year_path)):
            if fname.startswith('~$') or 'for copy' in fname.lower(): continue
            fpath = os.path.join(year_path, fname)
            if not os.path.isfile(fpath): continue
            
            subj = detect_subj(fname)
            if not subj: continue
            
            ext = os.path.splitext(fname)[1].lower()
            key = f'{year}|{subj}'
            if key not in data:
                data[key] = {'year': year, 'subject': subj, 'units': []}
                counts[key] = 0
            
            try:
                if ext == '.docx':
                    doc = docx.Document(fpath)
                    for table in doc.tables:
                        rows = extract_table_rows(table)
                        if subj in ('BM',):
                            units = process_bm_table(rows)
                        elif subj in ('BI',):
                            units = process_bi_table(rows)
                        elif subj in ('MT', 'SN', 'BC', 'PM', 'PI', 'PJ', 'PK', 'PSV', 'MZ', 'RBT', 'SJ'):
                            units = process_mt_table(rows)
                        else:
                            units = process_mt_table(rows)
                        data[key]['units'].extend(units)
                elif ext == '.xlsx':
                    wb = openpyxl.load_workbook(fpath, data_only=True)
                    ws = wb.active
                    rows = []
                    for row in ws.iter_rows(values_only=True):
                        cells = [clean(c) if c else '' for c in row]
                        if any(c for c in cells):
                            rows.append(cells)
                    units = process_mt_table(rows)
                    data[key]['units'].extend(units)
            except Exception as e:
                print(f"  Error {fname}: {e}", file=sys.stderr)
    
    # Deduplicate units per key
    for key in data:
        seen_titles = set()
        clean_units = []
        for u in data[key]['units']:
            t = u.get('title', '')[:50]
            if t in seen_titles: continue
            seen_titles.add(t)
            clean_units.append(u)
        data[key]['units'] = clean_units
        print(f"  {key}: {len(clean_units)} units")
    
    return data

def build_js(data):
    lines = []
    lines.append('// eRPH-PM Curriculum Data — Auto-generated from official RPT files')
    lines.append(f'// Generated: {datetime.now().isoformat()}')
    lines.append('')
    
    # SUBJECTS array
    lines.append('const SUBJECTS = [')
    subj_list = sorted(set(v['subject'] for v in data.values()), key=lambda s: list(SUBJECT_INFO.keys()).index(s) if s in SUBJECT_INFO else 99)
    for i, s in enumerate(subj_list):
        info = SUBJECT_INFO.get(s, {'name': s, 'code': s, 'color': '#666'})
        comma = ',' if i < len(subj_list) - 1 else ''
        lines.append(f"  {{ code: '{s}', name: '{info['name']}', color: '{info['color']}' }}{comma}")
    lines.append('];')
    lines.append('')
    
    # YEAR_GROUPS
    lines.append("const YEAR_GROUPS = ['Year 1','Year 2','Year 3','Year 4','Year 5','Year 6'];")
    lines.append('')
    
    # CURRICULUM_DATA
    lines.append('const CURRICULUM_DATA = {')
    
    keys = sorted(data.keys())
    for ki, key in enumerate(keys):
        entry = data[key]
        year = entry['year']
        subj = entry['subject']
        units = entry['units']
        
        if ki > 0: lines[-1] = lines[-1] + ','
        lines.append(f"  '{year}|{subj}': {{")
        lines.append(f"    year: '{year}',")
        lines.append(f"    subject: '{subj}',")
        lines.append(f"    subjectName: '{SUBJECT_INFO.get(subj, {}).get('name', subj)}',")
        lines.append("    units: [")
        
        for ui, u in enumerate(units):
            title = u.get('title', '').replace("'", "\\'").replace('"', "'")
            minggu = u.get('minggu', '').replace("'", "\\'").replace('"', "'")[:30]
            sk = u.get('sk', '').replace("'", "\\'").replace('"', "'")[:200]
            sp = u.get('sp', '').replace("'", "\\'").replace('"', "'")[:200]
            
            if ui > 0: lines[-1] = lines[-1] + ','
            lines.append(f"      {{ id: {ui+1}, title: '{title}', minggu: '{minggu}', sk: '{sk}', sp: '{sp}' }}")
        
        if not units:
            lines.append("      []")
        lines.append("    ]")
        lines.append("  }")
    
    lines.append('};')
    lines.append('')
    
    # TEXTBOOK_INFO
    lines.append('const TEXTBOOK_INFO = {')
    for ki, key in enumerate(keys):
        entry = data[key]
        year = entry['year']
        subj = entry['subject']
        name = SUBJECT_INFO.get(subj, {}).get('name', subj)
        comma = ',' if ki < len(keys) - 1 else ''
        lines.append(f"  '{year}|{subj}': {{ year: '{year}', subject: '{subj}', subjectName: '{name}', textbook: 'RPT {year} {name}', publisher: 'Ministry of Education Malaysia' }}{comma}")
    lines.append('};')
    
    lines.append('')
    lines.append('// TAKWIM 2026 - School calendar weeks')
    lines.append('const TAKWIM_2026 = [')
    weeks = [
        ('12/01/2026','16/01/2026',1,1,1),('19/01/2026','23/01/2026',2,1,1),('26/01/2026','30/01/2026',3,1,1),
        ('02/02/2026','06/02/2026',4,1,1),('09/02/2026','13/02/2026',5,1,1),('16/02/2026','20/02/2026',6,1,1),
        ('23/02/2026','27/02/2026',7,1,1),('02/03/2026','06/03/2026',8,1,1),('09/03/2026','13/03/2026',9,1,1),
        ('16/03/2026','20/03/2026',10,1,1),('21/03/2026','29/03/2026',0,1,2),('30/03/2026','03/04/2026',11,2,2),
        ('06/04/2026','10/04/2026',12,2,2),('13/04/2026','17/04/2026',13,2,2),('20/04/2026','24/04/2026',14,2,2),
        ('27/04/2026','01/05/2026',15,2,2),('04/05/2026','08/05/2026',16,2,2),('11/05/2026','15/05/2026',17,2,2),
        ('18/05/2026','22/05/2026',18,2,2),('25/05/2026','29/05/2026',19,2,2),('01/06/2026','05/06/2026',20,2,2),
        ('08/06/2026','12/06/2026',21,2,2),('15/06/2026','19/06/2026',22,2,2),('22/06/2026','26/06/2026',23,2,2),
        ('29/06/2026','03/07/2026',24,2,2),('06/07/2026','10/07/2026',25,2,2),('13/07/2026','17/07/2026',26,2,2),
        ('20/07/2026','24/07/2026',27,2,2),('27/07/2026','31/07/2026',28,2,2),('03/08/2026','07/08/2026',29,2,2),
        ('10/08/2026','14/08/2026',30,2,2),('17/08/2026','21/08/2026',31,2,2),('24/08/2026','28/08/2026',32,2,2),
        ('31/08/2026','04/09/2026',33,2,2),('07/09/2026','11/09/2026',34,2,2),('14/09/2026','18/09/2026',35,2,2),
        ('28/09/2026','02/10/2026',36,2,2),('05/10/2026','09/10/2026',37,2,2),('12/10/2026','16/10/2026',38,2,2),
        ('19/10/2026','23/10/2026',39,2,2),('26/10/2026','30/10/2026',40,2,2),('02/11/2026','06/11/2026',41,2,2),
        ('09/11/2026','13/11/2026',42,2,2),('16/11/2026','20/11/2026',43,2,2),('23/11/2026','27/11/2026',44,2,2),
        ('30/11/2026','04/12/2026',45,2,2),('07/12/2026','11/12/2026',46,2,2),('14/12/2026','18/12/2026',47,2,2),
    ]
    for i, (s, e, w, t, b) in enumerate(weeks):
        comma = ',' if i < len(weeks) - 1 else ''
        lines.append(f"  {{ start: '{s}', end: '{e}', minggu: {w}, term: {t}, block: {b} }}{comma}")
    lines.append('];')
    
    lines.append('')
    lines.append('// Default timetable template (can be customized)')
    lines.append('const DEFAULT_TIMETABLE = [')
    lines.append('];')
    
    lines.append('')
    lines.append('// Ensure backward compatibility')
    lines.append("if (typeof module !== 'undefined' && module.exports) { module.exports = { CURRICULUM_DATA, TEXTBOOK_INFO, SUBJECTS, YEAR_GROUPS, TAKWIM_2026, DEFAULT_TIMETABLE }; }")
    lines.append('')
    
    return '\n'.join(lines)

if __name__ == '__main__':
    print("Extracting curriculum data from RPT files...")
    data = extract_all()
    print(f"\nTotal: {len(data)} subject/year combinations")
    
    js = build_js(data)
    output_path = os.path.join(os.path.dirname(__file__), 'curriculum-data.js')
    with open(output_path, 'w', encoding='utf-8') as f:
        f.write(js)
    print(f"Written to {output_path}")
    print(f"Size: {os.path.getsize(output_path)} bytes")
